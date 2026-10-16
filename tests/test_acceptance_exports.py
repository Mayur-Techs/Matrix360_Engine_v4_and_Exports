"""Acceptance-level checks that do not require a network socket or real client data."""
from __future__ import annotations

import csv

from openpyxl import load_workbook

from matrix_engine.engine.core import run
from services.worker.exporter import create_exports


def write_csv(path, rows):
    with path.open("w", encoding="utf-8", newline="") as handle:
        csv.writer(handle).writerows(rows)


def test_three_clients_geography_dedup_and_export_reconciliation(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    write_csv(source / "fresh.csv", [
        ["Name", "Phone", "Email", "City", "State", "Company"],
        ["Industrial", "09876543210", "shared@example.test", "Howrah", "West Bengal", "Steel Works"],
        ["Residential A", "9876543210", "shared@example.test", "Chandannagar", "West Bengal", ""],
        ["Residential B", "9876543211", "", "Kalyani", "West Bengal", ""],
        ["Outside", "9876543212", "", "Pune", "Maharashtra", "Steel Works"],
        ["Unknown", "9876543213", "", "", "", ""],
    ])
    workspace = tmp_path / "workspace"
    report = run(source, workspace)
    counts, artifacts = create_exports(workspace, tmp_path / "exports", "job-acceptance", "tenant-acceptance", report["version"])

    assert report["reconciled"] is True
    assert report["reconciliation"]["input_rows_processed"] == 5
    assert report["reconciliation"]["geography_counts"] == {"GEO_UNKNOWN": 1, "OUTSIDE_WB": 1, "WB_CONFIRMED": 3}
    assert report["reconciliation"]["deduplication_review_groups"] == 1
    assert report["reconciliation"]["automatic_merges"] == 0
    assert counts["Lakhotia"] == counts["Shree_Heights"] == counts["Vrindavan_Greens"] == 1
    assert counts["Outside_WB"] == counts["Review_Required"] == 1
    assert counts["Unique_export_candidates"] == 3
    assert {artifact["name"] for artifact in artifacts} >= {"Lakhotia.csv", "Shree_Heights.csv", "Vrindavan_Greens.csv", "Matrix360_Combined.xlsx"}

    with (tmp_path / "exports" / "Lakhotia.csv").open(encoding="utf-8", newline="") as handle:
        output = list(csv.DictReader(handle))
    assert output[0]["phone_source"] == "09876543210"

    workbook = load_workbook(tmp_path / "exports" / "Matrix360_Combined.xlsx", read_only=True)
    assert {"Summary", "Lakhotia", "Shree_Heights", "Vrindavan_Greens", "WB_Confirmed", "WB_Possible", "Review_Required"} <= set(workbook.sheetnames)
    workbook.close()
