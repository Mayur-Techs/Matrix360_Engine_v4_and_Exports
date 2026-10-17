"""Regression checks for persisted client and district export filters."""
from __future__ import annotations

import csv
import json

from openpyxl import load_workbook

from matrix_engine.engine.core import run
from services.worker.exporter import create_exports


def write_csv(path, rows):
    with path.open("w", encoding="utf-8", newline="") as handle:
        csv.writer(handle).writerows(rows)


def test_client_and_district_filters_limit_only_client_exports(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    write_csv(source / "campaign.csv", [
        ["Name", "Phone", "City", "State", "Company"],
        ["Howrah industrial", "09876543210", "Howrah", "West Bengal", "Steel Works"],
        ["Kolkata industrial", "09876543211", "Kolkata", "West Bengal", "Steel Works"],
        ["Howrah residential", "09876543212", "Howrah", "West Bengal", ""],
    ])
    workspace = tmp_path / "workspace"
    report = run(source, workspace)

    counts, artifacts = create_exports(
        workspace,
        tmp_path / "exports",
        "job-filtered",
        "tenant-filtered",
        report["version"],
        client_profiles=("lakhotia",),
        districts=("Howrah",),
        dataset_name="campaign.csv",
    )

    assert counts["Lakhotia"] == 1
    assert counts["Unique_export_candidates"] == 1
    assert "Shree_Heights" not in counts
    assert {artifact["name"] for artifact in artifacts} >= {"Lakhotia.csv", "Matrix360_Combined.xlsx", "job_manifest.v1.json"}
    assert not (tmp_path / "exports" / "Shree_Heights.csv").exists()

    with (tmp_path / "exports" / "Lakhotia.csv").open(encoding="utf-8", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert [row["name"] for row in rows] == ["Howrah industrial"]
    assert rows[0]["phone_source"] == "09876543210"

    manifest = json.loads((tmp_path / "exports" / "job_manifest.v1.json").read_text(encoding="utf-8"))
    assert manifest["dataset_name"] == "campaign.csv"
    assert manifest["filters"] == {"client_profiles": ["lakhotia"], "districts": ["Howrah"]}

    workbook = load_workbook(tmp_path / "exports" / "Matrix360_Combined.xlsx", read_only=True)
    assert "Lakhotia" in workbook.sheetnames
    assert "Shree_Heights" not in workbook.sheetnames
    summary = dict(workbook["Summary"].iter_rows(min_row=2, values_only=True))
    assert summary["Dataset"] == "campaign.csv"
    assert summary["District filter"] == "Howrah"
    workbook.close()
