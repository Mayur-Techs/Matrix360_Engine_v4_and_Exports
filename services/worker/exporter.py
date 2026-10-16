"""Bounded-memory export generation for completed local jobs."""
from __future__ import annotations

import csv
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable, Iterable

from openpyxl import Workbook

PROJECTS = {
    "Lakhotia": "lakhotia_class",
    "Shree_Heights": "shree_class",
    "Vrindavan_Greens": "vrindavan_class",
}
HEADERS = [
    "source_record_id", "source_file", "source_sheet", "source_row", "name", "company",
    "phone_source", "phone_normalized", "email_normalized", "city", "district", "state", "geo_status", "geo_evidence", "lakhotia_class",
    "shree_class", "vrindavan_class", "reason_codes", "rationale", "quality_flags", "rule_version",
]
MAX_XLSX_ROWS = 1_048_576


def safe_cell(value: object) -> str:
    value = "" if value is None else str(value)
    return "'" + value if value[:1] in ("=", "+", "-", "@") else value


def source_rows(path: Path) -> Iterable[dict[str, str]]:
    with path.open("r", encoding="utf-8", newline="") as handle:
        yield from csv.DictReader(handle)


def selected(row: dict[str, str], project_column: str) -> bool:
    return row.get("geo_status") == "WB_CONFIRMED" and row.get(project_column) in {"Q1", "Q2", "P1", "P2"}


def review(row: dict[str, str]) -> bool:
    return row.get("geo_status") in {"WB_POSSIBLE", "GEO_UNKNOWN", "REVIEW_REQUIRED"}


def write_csv(source: Path, destination: Path, predicate: Callable[[dict[str, str]], bool]) -> int:
    count = 0
    with destination.open("w", encoding="utf-8", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=HEADERS, extrasaction="ignore")
        writer.writeheader()
        for row in source_rows(source):
            if predicate(row):
                writer.writerow({header: safe_cell(row.get(header, "")) for header in HEADERS})
                count += 1
    return count


def add_partitioned_sheet(workbook: Workbook, title: str, source: Path, predicate: Callable[[dict[str, str]], bool]) -> int:
    total = 0
    sheet = None
    row_in_sheet = MAX_XLSX_ROWS
    partition = 0
    for row in source_rows(source):
        if not predicate(row):
            continue
        if row_in_sheet >= MAX_XLSX_ROWS:
            partition += 1
            suffix = "" if partition == 1 else f"_{partition}"
            sheet = workbook.create_sheet(f"{title}{suffix}")
            sheet.append(HEADERS)
            row_in_sheet = 1
        assert sheet is not None
        sheet.append([safe_cell(row.get(header, "")) for header in HEADERS])
        row_in_sheet += 1
        total += 1
    if sheet is None:
        sheet = workbook.create_sheet(title)
        sheet.append(HEADERS)
    return total


def digest(path: Path) -> str:
    hash_ = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            hash_.update(block)
    return hash_.hexdigest()


def create_exports(workspace: Path, export_dir: Path, job_id: str, tenant_id: str, processing_version: str) -> tuple[dict[str, int], list[dict[str, str | int]]]:
    source = workspace / "all_records.csv"
    if not source.exists():
        raise FileNotFoundError("Worker output all_records.csv is missing")
    export_dir.mkdir(parents=True, exist_ok=True)
    counts: dict[str, int] = {}
    artifacts: list[dict[str, str | int]] = []
    for name, column in PROJECTS.items():
        file = export_dir / f"{name}.csv"
        counts[name] = write_csv(source, file, lambda row, c=column: selected(row, c))
        artifacts.append({"name": file.name, "path": str(file), "content_type": "text/csv", "size_bytes": file.stat().st_size, "sha256": digest(file)})

    outside = write_csv(source, export_dir / "Outside_WB_review.csv", lambda row: row.get("geo_status") == "OUTSIDE_WB")
    review_count = write_csv(source, export_dir / "Review_Required.csv", review)
    counts["Outside_WB"] = outside
    counts["Review_Required"] = review_count

    workbook_path = export_dir / "Matrix360_Combined.xlsx"
    workbook = Workbook(write_only=True)
    summary = workbook.create_sheet("Summary")
    summary.append(["Metric", "Value"])
    summary.append(["Job ID", job_id])
    summary.append(["Tenant", tenant_id])
    summary.append(["Generated at", datetime.now(timezone.utc).isoformat()])
    summary.append(["Processing version", processing_version])
    summary.append(["Engine mode", "Matrix Engine V4 local rulebase"])
    summary.append(["Permission notice", "Exports are not consent to contact."])
    for name, column in PROJECTS.items():
        counts[name] = add_partitioned_sheet(workbook, name, source, lambda row, c=column: selected(row, c))
        summary.append([name, counts[name]])
    counts["WB_Confirmed"] = add_partitioned_sheet(workbook, "WB_Confirmed", source, lambda row: row.get("geo_status") == "WB_CONFIRMED")
    counts["WB_Possible"] = add_partitioned_sheet(workbook, "WB_Possible", source, lambda row: row.get("geo_status") == "WB_POSSIBLE")
    add_partitioned_sheet(workbook, "Review_Required", source, review)
    counts["Unique_export_candidates"] = sum(1 for row in source_rows(source) if any(selected(row, column) for column in PROJECTS.values()))
    for name in ("WB_Confirmed", "WB_Possible", "Outside_WB", "Review_Required", "Unique_export_candidates"):
        summary.append([name.replace("_", " "), counts[name]])
    workbook.save(workbook_path)
    artifacts.append({"name": workbook_path.name, "path": str(workbook_path), "content_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "size_bytes": workbook_path.stat().st_size, "sha256": digest(workbook_path)})

    manifest_path = export_dir / "job_manifest.v1.json"
    manifest = {
        "job_id": job_id,
        "tenant_id": tenant_id,
        "status": "completed",
        "processing_version": processing_version,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "counts": counts,
        "artifacts": artifacts,
    }
    manifest_path.write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    artifacts.append({"name": manifest_path.name, "path": str(manifest_path), "content_type": "application/json", "size_bytes": manifest_path.stat().st_size, "sha256": digest(manifest_path)})
    return counts, artifacts

