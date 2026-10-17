"""Authenticated local Matrix 360 API and same-origin console."""
from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import uuid
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Annotated, Any

import duckdb
import openpyxl
from fastapi import Depends, FastAPI, File, Form, HTTPException, Query, Request, UploadFile, status
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from . import store
from .security import User, authenticate, current_user, issue, verify
from .settings import COOKIE_SECURE, LOCAL_PASSWORD, MAX_UPLOAD_BYTES, RUNTIME_ROOT, validate_deployment_settings

ROOT = Path(__file__).resolve().parents[2]
WEB_ROOT = ROOT / "apps" / "web"
ALLOWED_SUFFIXES = {".csv", ".tsv", ".xlsx", ".xls", ".xlsb", ".parquet", ".jsonl", ".ndjson"}
FILTERABLE_RESULT_FIELDS = {
    "name": {"label": "Name", "group": "Person and company"},
    "company": {"label": "Company", "group": "Person and company"},
    "job_title": {"label": "Job title / designation", "group": "Person and company"},
    "department": {"label": "Department", "group": "Person and company"},
    "sector": {"label": "Industry / sector", "group": "Person and company"},
    "state": {"label": "State", "group": "Location"},
    "district": {"label": "District", "group": "Location"},
    "city": {"label": "City", "group": "Location"},
    "postal_code": {"label": "PIN / postal code", "group": "Location"},
    "dataset_category": {"label": "Dataset category", "group": "Source and data quality"},
    "source_file": {"label": "Source file", "group": "Source and data quality"},
    "source_sheet": {"label": "Source sheet", "group": "Source and data quality"},
    "quality_flags": {"label": "Data-quality flag", "group": "Source and data quality"},
    "vehicle_registration": {"label": "Vehicle registration", "group": "Vehicle (observed data only)"},
    "vehicle_make": {"label": "Vehicle make / brand", "group": "Vehicle (observed data only)"},
    "vehicle_model": {"label": "Vehicle model", "group": "Vehicle (observed data only)"},
    "vehicle_variant": {"label": "Vehicle variant", "group": "Vehicle (observed data only)"},
    "vehicle_fuel": {"label": "Vehicle fuel", "group": "Vehicle (observed data only)"},
    "vehicle_year": {"label": "Vehicle year", "group": "Vehicle (observed data only)"},
}
ADVANCED_FILTER_OPERATORS = {"contains", "not_contains", "is_any_of", "is_none_of", "is_present", "is_empty"}


class LoginRequest(BaseModel):
    username: str
    password: str


class JobRequest(BaseModel):
    dataset_id: str
    idempotency_key: str = Field(min_length=8, max_length=128)
    districts: list[str] = Field(default_factory=list, max_length=23)
    client_profiles: list[str] = Field(default_factory=lambda: ["lakhotia", "shree", "vrindavan"], max_length=3)
    ai_review_enabled: bool = False


@asynccontextmanager
async def lifespan(_: FastAPI):
    validate_deployment_settings()
    store.init()
    RUNTIME_ROOT.mkdir(parents=True, exist_ok=True)
    yield


app = FastAPI(title="Matrix 360 Audience Intelligence API", version="1.0.0", lifespan=lifespan)


@app.get("/healthz", tags=["system"])
def healthz() -> dict[str, Any]:
    return {
        "status": "ok",
        "mode": "local-single-user",
        "engine": "matrix-engine-v4-rulebase-local",
        "ai_configured": bool(__import__("os").environ.get("OPENAI_API_KEY")),
        "credential_warning": LOCAL_PASSWORD == "change-me-before-sharing",
    }


@app.post("/api/v1/auth/login", tags=["auth"])
def login(body: LoginRequest) -> JSONResponse:
    user = authenticate(body.username, body.password)
    if not user:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Invalid credentials")
    response = JSONResponse({"username": user.username, "tenant_id": user.tenant_id, "role": user.role})
    response.set_cookie("matrix_session", issue(user), httponly=True, samesite="lax", secure=COOKIE_SECURE, max_age=8 * 60 * 60)
    store.audit(user.tenant_id, user.username, "auth.login", user.username)
    return response


@app.post("/api/v1/auth/logout", tags=["auth"])
def logout(user: Annotated[User, Depends(current_user)]) -> JSONResponse:
    response = JSONResponse({"ok": True})
    response.delete_cookie("matrix_session")
    store.audit(user.tenant_id, user.username, "auth.logout", user.username)
    return response


@app.get("/api/v1/auth/me", tags=["auth"])
def me(user: Annotated[User, Depends(current_user)]) -> User:
    return user


@app.get("/api/v1/auth/session", tags=["auth"])
def session_status(request: Request) -> dict[str, Any]:
    """Anonymous-safe session probe so the login screen never emits a 401 console error."""
    user = verify(request.cookies.get("matrix_session"))
    return {"authenticated": bool(user), "user": user if user else None}


def safe_name(filename: str | None) -> str:
    name = Path(filename or "").name.strip()
    if not name or name.startswith(".") or name.lower() in {"thumbs.db", "desktop.ini"}:
        raise HTTPException(400, "Unsafe file name")
    if Path(name).suffix.lower() not in ALLOWED_SUFFIXES:
        raise HTTPException(415, f"Unsupported file format: {Path(name).suffix or 'unknown'}")
    return re.sub(r"[^A-Za-z0-9._ -]", "_", name)


def safe_display_path(value: str, fallback: str) -> str:
    candidate = Path(value.replace("\\", "/"))
    if candidate.is_absolute() or ".." in candidate.parts or any(part.startswith(".") for part in candidate.parts):
        return fallback
    return "/".join(re.sub(r"[^A-Za-z0-9._ -]", "_", part) for part in candidate.parts if part) or fallback


def dataset_display_name(conn: Any, dataset_id: str) -> str:
    """Return an operator-safe label for a dataset without exposing local paths."""
    rows = conn.execute(
        "SELECT original_name FROM dataset_files WHERE dataset_id=? ORDER BY relative_path",
        (dataset_id,),
    ).fetchall()
    names = [str(row["original_name"]) for row in rows]
    return ", ".join(names) if names else dataset_id


def imported_headers(conn: Any, dataset_id: str) -> list[str]:
    """Return unique uploaded headers for a filter builder without returning row data."""
    row = conn.execute("SELECT schema_json FROM datasets WHERE id=?", (dataset_id,)).fetchone()
    if not row:
        return []
    try:
        files = json.loads(row["schema_json"]).get("files", [])
    except (TypeError, json.JSONDecodeError):
        return []
    values = {
        str(header).strip()
        for item in files
        for header in item.get("inspection", {}).get("filter_headers", item.get("inspection", {}).get("headers", []))
        if str(header).strip() and len(str(header).strip()) <= 160
    }
    return sorted(values, key=str.lower)


def parse_advanced_filters(value: str, source_headers: list[str]) -> list[dict[str, Any]]:
    """Validate schema-aware filter rules before they are translated to SQL."""
    if not value:
        return []
    try:
        filters = json.loads(value)
    except json.JSONDecodeError as exc:
        raise HTTPException(422, "Advanced filters must be valid JSON") from exc
    if not isinstance(filters, list) or len(filters) > 12:
        raise HTTPException(422, "Use between zero and twelve advanced filters")
    allowed_source_fields = set(source_headers)
    normalized: list[dict[str, Any]] = []
    for item in filters:
        if not isinstance(item, dict):
            raise HTTPException(422, "Each advanced filter must be an object")
        field = str(item.get("field", ""))
        operator = str(item.get("operator", ""))
        values = item.get("values", [])
        if field not in FILTERABLE_RESULT_FIELDS and field not in allowed_source_fields:
            raise HTTPException(422, "Unsupported advanced filter field")
        if operator not in ADVANCED_FILTER_OPERATORS:
            raise HTTPException(422, "Unsupported advanced filter operator")
        if not isinstance(values, list) or len(values) > 25:
            raise HTTPException(422, "Each advanced filter may have up to 25 values")
        safe_values = [str(candidate).strip()[:120] for candidate in values if str(candidate).strip()]
        if operator not in {"is_present", "is_empty"} and not safe_values:
            raise HTTPException(422, "This filter needs at least one value")
        normalized.append({"field": field, "operator": operator, "values": safe_values})
    return normalized


def advanced_filter_sql(filters: list[dict[str, Any]]) -> tuple[list[str], list[str]]:
    """Build parameterized DuckDB predicates for allow-listed result/source fields."""
    clauses: list[str] = []
    params: list[str] = []
    for filter_ in filters:
        field = filter_["field"]
        if field in FILTERABLE_RESULT_FIELDS:
            expression = field
        else:
            escaped = field.replace("\\", "\\\\").replace('"', '\\"')
            json_path = f'$."{escaped}"'.replace("'", "''")
            expression = f"json_extract_string(original_fields, '{json_path}')"
        normalized = f"lower(coalesce({expression}, ''))"
        values = [value.lower() for value in filter_["values"]]
        operator = filter_["operator"]
        if operator == "is_present":
            clauses.append(f"{normalized} <> ''")
        elif operator == "is_empty":
            clauses.append(f"{normalized} = ''")
        elif operator == "contains":
            clauses.append("(" + " OR ".join(f"{normalized} LIKE ?" for _ in values) + ")")
            params.extend([f"%{value}%" for value in values])
        elif operator == "not_contains":
            clauses.append("(" + " AND ".join(f"{normalized} NOT LIKE ?" for _ in values) + ")")
            params.extend([f"%{value}%" for value in values])
        elif operator == "is_any_of":
            clauses.append(f"{normalized} IN (" + ",".join("?" for _ in values) + ")")
            params.extend(values)
        else:  # is_none_of
            clauses.append(f"{normalized} NOT IN (" + ",".join("?" for _ in values) + ")")
            params.extend(values)
    return clauses, params


def inspect_upload(path: Path) -> dict[str, Any]:
    sample: list[list[str]] = []
    headers: list[str] = []
    filter_headers: list[str] = []
    try:
        if path.suffix.lower() in {".csv", ".tsv"}:
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                reader = csv.reader(handle, delimiter="\t" if path.suffix.lower() == ".tsv" else ",")
                headers = [str(value) for value in next(reader, [])]
                sample = [[str(value) for value in row] for _, row in zip(range(5), reader)]
                filter_headers = headers
        elif path.suffix.lower() == ".parquet":
            relation = duckdb.read_parquet(str(path))
            headers = list(relation.columns)
            sample = [["" if value is None else str(value) for value in row] for row in relation.limit(5).fetchall()]
            filter_headers = headers
        elif path.suffix.lower() in {".xlsx", ".xlsm"}:
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                sheet = workbook.active
                rows = sheet.iter_rows(values_only=True)
                headers = ["" if value is None else str(value) for value in next(rows, ())]
                sample = [["" if value is None else str(value) for value in row] for _, row in zip(range(5), rows)]
                filter_headers = [
                    "" if value is None else str(value)
                    for candidate in workbook.worksheets
                    for value in next(candidate.iter_rows(values_only=True), ())
                ]
            finally:
                workbook.close()
        else:
            headers = ["Schema inspection begins when processing starts"]
            filter_headers = headers
    except Exception as exc:
        headers = ["Inspection warning"]
        sample = [[type(exc).__name__]]
        filter_headers = headers
    return {"headers": headers, "filter_headers": filter_headers, "sample": sample}


@app.post("/api/v1/datasets", tags=["datasets"])
async def upload_datasets(
    request: Request,
    files: Annotated[list[UploadFile], File(description="One or more supported data files")],
    relative_paths: Annotated[str | None, Form()] = None,
    user: User = Depends(current_user),
) -> dict[str, Any]:
    if not files:
        raise HTTPException(400, "Choose at least one file")
    try:
        requested_paths = json.loads(relative_paths) if relative_paths else []
    except json.JSONDecodeError:
        raise HTTPException(400, "relative_paths must be JSON")
    if requested_paths and (not isinstance(requested_paths, list) or len(requested_paths) != len(files)):
        raise HTTPException(400, "relative_paths must match uploaded files")
    dataset_id = str(uuid.uuid4())
    base = RUNTIME_ROOT / "uploads" / user.tenant_id / dataset_id
    base.mkdir(parents=True, exist_ok=False)
    records: list[dict[str, Any]] = []
    try:
        for index, upload in enumerate(files):
            name = safe_name(upload.filename)
            destination = base / f"{index:04d}_{name}"
            sha = hashlib.sha256()
            size = 0
            with destination.open("wb") as output:
                while block := await upload.read(1024 * 1024):
                    size += len(block)
                    if size > MAX_UPLOAD_BYTES:
                        raise HTTPException(413, f"{name} exceeds local upload limit")
                    sha.update(block)
                    output.write(block)
            display_path = safe_display_path(str(requested_paths[index]) if requested_paths else name, name)
            records.append({"name": name, "stored_path": str(destination), "relative_path": display_path, "size_bytes": size, "sha256": sha.hexdigest(), "inspection": inspect_upload(destination)})
    except Exception:
        shutil.rmtree(base, ignore_errors=True)
        raise
    schema = {"files": records}
    with store.connect() as conn:
        conn.execute(
            "INSERT INTO datasets(id,tenant_id,uploaded_at,schema_json) VALUES(?,?,?,?)",
            (dataset_id, user.tenant_id, store.now(), json.dumps(schema)),
        )
        for record in records:
            conn.execute(
                "INSERT INTO dataset_files(id,dataset_id,original_name,stored_path,relative_path,size_bytes,sha256) VALUES(?,?,?,?,?,?,?)",
                (str(uuid.uuid4()), dataset_id, record["name"], record["stored_path"], record["relative_path"], record["size_bytes"], record["sha256"]),
            )
    store.audit(user.tenant_id, user.username, "dataset.uploaded", dataset_id, {"file_count": len(records), "bytes": sum(item["size_bytes"] for item in records)})
    return {"dataset_id": dataset_id, "files": records, "mode": "local-streamed"}


@app.get("/api/v1/datasets", tags=["datasets"])
def list_datasets(user: Annotated[User, Depends(current_user)]) -> dict[str, Any]:
    with store.connect() as conn:
        rows = conn.execute(
            "SELECT d.id,d.uploaded_at,f.original_name,f.relative_path,f.size_bytes,f.sha256 FROM datasets d JOIN dataset_files f ON f.dataset_id=d.id WHERE d.tenant_id=? ORDER BY d.uploaded_at DESC,f.relative_path",
            (user.tenant_id,),
        ).fetchall()
    grouped: dict[str, dict[str, Any]] = {}
    for row in rows:
        item = dict(row)
        grouped.setdefault(item["id"], {"id": item.pop("id"), "uploaded_at": item.pop("uploaded_at"), "files": []})["files"].append(item)
    return {"datasets": list(grouped.values())}


@app.post("/api/v1/jobs", tags=["jobs"])
def create_job(body: JobRequest, user: Annotated[User, Depends(current_user)]) -> dict[str, Any]:
    known_profiles = {"lakhotia", "shree", "vrindavan"}
    profiles = sorted(set(body.client_profiles))
    if not profiles:
        raise HTTPException(422, "Choose at least one client profile")
    if not set(profiles).issubset(known_profiles):
        raise HTTPException(422, "Unknown client profile")
    districts = sorted({district.strip() for district in body.districts if district.strip()})
    if body.ai_review_enabled:
        # This endpoint deliberately does not enable external transfer: an owner must opt in through deployment config too.
        from .settings import AI_ALLOW_PII, AI_ENABLED, AI_MAX_CALLS
        if not (AI_ENABLED and AI_ALLOW_PII and AI_MAX_CALLS > 0):
            raise HTTPException(409, "AI review is deployment-disabled pending explicit data-transfer and spend approval")
    config = {"districts": districts, "client_profiles": profiles, "ai_review_enabled": body.ai_review_enabled}
    with store.connect() as conn:
        dataset = conn.execute("SELECT 1 FROM datasets WHERE id=? AND tenant_id=?", (body.dataset_id, user.tenant_id)).fetchone()
        if not dataset:
            raise HTTPException(404, "Dataset not found")
        existing = conn.execute("SELECT * FROM jobs WHERE tenant_id=? AND idempotency_key=?", (user.tenant_id, body.idempotency_key)).fetchone()
        if existing:
            return {"job": dict(existing), "idempotent_replay": True}
        job_id = str(uuid.uuid4())
        workspace = RUNTIME_ROOT / "workspaces" / user.tenant_id / job_id
        conn.execute(
            "INSERT INTO jobs(id,tenant_id,dataset_id,idempotency_key,status,created_at,workspace_path,config_json) VALUES(?,?,?,?,?,?,?,?)",
            (job_id, user.tenant_id, body.dataset_id, body.idempotency_key, "queued", store.now(), str(workspace), json.dumps(config, sort_keys=True)),
        )
    store.audit(user.tenant_id, user.username, "job.queued", job_id, {"dataset_id": body.dataset_id, "config": config})
    return {"job_id": job_id, "status": "queued", "worker": "Secure background worker will process this job automatically."}


@app.get("/api/v1/jobs", tags=["jobs"])
def list_jobs(user: Annotated[User, Depends(current_user)]) -> dict[str, Any]:
    with store.connect() as conn:
        rows = conn.execute("SELECT id,dataset_id,status,created_at,started_at,completed_at,error_code,error_detail,progress_phase,progress_percent,report_json FROM jobs WHERE tenant_id=? ORDER BY created_at DESC", (user.tenant_id,)).fetchall()
        dataset_names = {row["dataset_id"]: dataset_display_name(conn, row["dataset_id"]) for row in rows}
    result = []
    for row in rows:
        item = dict(row)
        item["report"] = json.loads(item.pop("report_json") or "{}")
        item["dataset_name"] = dataset_names[item["dataset_id"]]
        result.append(item)
    return {"jobs": result}


@app.get("/api/v1/jobs/{job_id}", tags=["jobs"])
def get_job(job_id: str, user: Annotated[User, Depends(current_user)]) -> dict[str, Any]:
    with store.connect() as conn:
        row = conn.execute("SELECT * FROM jobs WHERE id=? AND tenant_id=?", (job_id, user.tenant_id)).fetchone()
        dataset_name = dataset_display_name(conn, row["dataset_id"]) if row else ""
    if not row:
        raise HTTPException(404, "Job not found")
    item = dict(row)
    item["report"] = json.loads(item.pop("report_json") or "{}")
    item["dataset_name"] = dataset_name
    return item


@app.get("/api/v1/jobs/{job_id}/results", tags=["results"])
def results(job_id: str, user: Annotated[User, Depends(current_user)], offset: int = 0, limit: int = 50, search: str = "", geo_status: str = "", geo_statuses: Annotated[list[str], Query()] = [], districts: Annotated[list[str], Query()] = [], client_profiles: Annotated[list[str], Query()] = [], advanced_filters: str = "") -> dict[str, Any]:
    limit = max(1, min(limit, 200))
    with store.connect() as conn:
        job = conn.execute("SELECT workspace_path,status,dataset_id FROM jobs WHERE id=? AND tenant_id=?", (job_id, user.tenant_id)).fetchone()
        source_headers = imported_headers(conn, job["dataset_id"]) if job else []
    if not job:
        raise HTTPException(404, "Job not found")
    source = Path(job["workspace_path"]) / "all_records.csv"
    if not source.exists():
        return {"rows": [], "total": 0, "status": job["status"]}
    with source.open("r", encoding="utf-8", newline="") as handle:
        result_headers = set(next(csv.reader(handle), []))
    clauses, params = [], []
    if search:
        clauses.append("(name ILIKE ? OR company ILIKE ? OR city ILIKE ? OR email_normalized ILIKE ?)")
        params.extend([f"%{search}%"] * 4)
    known_geo = {"WB_CONFIRMED", "WB_POSSIBLE", "OUTSIDE_WB", "GEO_UNKNOWN", "REVIEW_REQUIRED"}
    selected_geo = [value for value in [geo_status, *geo_statuses] if value in known_geo]
    if selected_geo:
        clauses.append("geo_status IN (" + ",".join("?" for _ in selected_geo) + ")")
        params.extend(selected_geo)
    if districts:
        clauses.append("district IN (" + ",".join("?" for _ in districts) + ")")
        params.extend(districts)
    column_by_profile = {"lakhotia": "lakhotia_class", "shree": "shree_class", "vrindavan": "vrindavan_class"}
    selected_profiles = [column_by_profile[value] for value in client_profiles if value in column_by_profile]
    if selected_profiles:
        clauses.append("(" + " OR ".join(f"{column} IN ('Q1','Q2','P1','P2')" for column in selected_profiles) + ")")
    parsed_advanced_filters = parse_advanced_filters(advanced_filters, source_headers)
    if any(item["field"] in FILTERABLE_RESULT_FIELDS and item["field"] not in result_headers for item in parsed_advanced_filters):
        raise HTTPException(422, "This processed job does not contain one or more selected filter fields")
    advanced_clauses, advanced_params = advanced_filter_sql(parsed_advanced_filters)
    clauses.extend(advanced_clauses)
    params.extend(advanced_params)
    where = " WHERE " + " AND ".join(clauses) if clauses else ""
    escaped = str(source).replace("'", "''")
    with duckdb.connect(":memory:") as conn:
        total = conn.execute(f"SELECT count(*) FROM read_csv_auto('{escaped}', all_varchar=true){where}", params).fetchone()[0]
        cursor = conn.execute(f"SELECT * FROM read_csv_auto('{escaped}', all_varchar=true){where} LIMIT ? OFFSET ?", [*params, limit, max(offset, 0)])
        columns = [description[0] for description in cursor.description]
        rows = [dict(zip(columns, values)) for values in cursor.fetchall()]
    return {"rows": rows, "total": total, "offset": offset, "limit": limit, "status": job["status"]}


@app.get("/api/v1/jobs/{job_id}/filter-metadata", tags=["results"])
def filter_metadata(job_id: str, user: Annotated[User, Depends(current_user)]) -> dict[str, Any]:
    """Expose only field availability/counts, never values, for schema-aware filters."""
    with store.connect() as conn:
        job = conn.execute("SELECT workspace_path,dataset_id,status FROM jobs WHERE id=? AND tenant_id=?", (job_id, user.tenant_id)).fetchone()
        source_headers = imported_headers(conn, job["dataset_id"]) if job else []
    if not job:
        raise HTTPException(404, "Job not found")
    source = Path(job["workspace_path"]) / "all_records.csv"
    supported = [{"id": key, **details, "available_records": 0} for key, details in FILTERABLE_RESULT_FIELDS.items()]
    if source.exists():
        with source.open("r", encoding="utf-8", newline="") as handle:
            result_headers = set((csv.reader(handle).__next__() if source.stat().st_size else []))
        available_columns = [field for field in FILTERABLE_RESULT_FIELDS if field in result_headers]
        if not available_columns:
            return {"status": job["status"], "operators": sorted(ADVANCED_FILTER_OPERATORS), "fields": supported, "imported_fields": source_headers, "vehicle_notice": "Vehicle fields are observed source attributes only. They do not establish residence, wealth, purchase intent, consent, or eligibility."}
        escaped = str(source).replace("'", "''")
        expressions = [f"count(*) FILTER (WHERE coalesce({field}, '') <> '') AS {field}" for field in available_columns]
        with duckdb.connect(":memory:") as conn:
            row = conn.execute(f"SELECT count(*) AS total_records,{','.join(expressions)} FROM read_csv_auto('{escaped}', all_varchar=true)").fetchone()
        counts = dict(zip(["total_records", *available_columns], row))
        for item in supported:
            item["available_records"] = int(counts.get(item["id"], 0) or 0)
    return {
        "status": job["status"],
        "operators": ["contains", "not_contains", "is_any_of", "is_none_of", "is_present", "is_empty"],
        "fields": supported,
        "imported_fields": source_headers,
        "vehicle_notice": "Vehicle fields are observed source attributes only. They do not establish residence, wealth, purchase intent, consent, or eligibility.",
    }


@app.get("/api/v1/jobs/{job_id}/exports", tags=["exports"])
def list_exports(job_id: str, user: Annotated[User, Depends(current_user)]) -> dict[str, Any]:
    with store.connect() as conn:
        job = conn.execute(
            "SELECT dataset_id,status,completed_at,config_json,report_json FROM jobs WHERE id=? AND tenant_id=?",
            (job_id, user.tenant_id),
        ).fetchone()
        rows = conn.execute("SELECT id,name,size_bytes,sha256,created_at,content_type FROM exports WHERE job_id=? AND tenant_id=? ORDER BY name", (job_id, user.tenant_id)).fetchall()
        dataset_name = dataset_display_name(conn, job["dataset_id"]) if job else ""
    if not job:
        raise HTTPException(404, "Job not found")
    config = json.loads(job["config_json"] or "{}")
    report = json.loads(job["report_json"] or "{}")
    return {
        "exports": [dict(row) for row in rows],
        "dataset_name": report.get("dataset_name") or dataset_name,
        "status": job["status"],
        "completed_at": job["completed_at"],
        "filters": report.get("export_filters") or {"client_profiles": config.get("client_profiles", []), "districts": config.get("districts", [])},
    }


@app.get("/api/v1/jobs/{job_id}/log/download", tags=["jobs"])
def download_job_log(job_id: str, user: Annotated[User, Depends(current_user)]) -> JSONResponse:
    """Download job metadata and audit events without exporting any source records."""
    with store.connect() as conn:
        job = conn.execute(
            "SELECT id,dataset_id,status,created_at,started_at,completed_at,error_code,error_detail,progress_phase,progress_percent,report_json,config_json FROM jobs WHERE id=? AND tenant_id=?",
            (job_id, user.tenant_id),
        ).fetchone()
        events = conn.execute(
            "SELECT actor,action,created_at,details_json FROM audit_events WHERE tenant_id=? AND entity_id=? ORDER BY id",
            (user.tenant_id, job_id),
        ).fetchall()
        dataset_name = dataset_display_name(conn, job["dataset_id"]) if job else ""
    if not job:
        raise HTTPException(404, "Job not found")
    payload = dict(job)
    payload["dataset_name"] = dataset_name
    payload["report"] = json.loads(payload.pop("report_json") or "{}")
    payload["configuration"] = json.loads(payload.pop("config_json") or "{}")
    payload["events"] = []
    for event in events:
        item = dict(event)
        item["details"] = json.loads(item.pop("details_json") or "{}")
        payload["events"].append(item)
    return JSONResponse(
        {"log_version": "matrix360-job-log-v1", "job": payload},
        headers={"Content-Disposition": f'attachment; filename="matrix360-job-{job_id[:8]}-log.json"'},
    )


@app.get("/api/v1/exports/{export_id}/download", tags=["exports"])
def download_export(export_id: str, user: Annotated[User, Depends(current_user)]) -> FileResponse:
    with store.connect() as conn:
        row = conn.execute("SELECT * FROM exports WHERE id=? AND tenant_id=?", (export_id, user.tenant_id)).fetchone()
    if not row or not Path(row["path"]).is_file():
        raise HTTPException(404, "Export not found")
    store.audit(user.tenant_id, user.username, "export.downloaded", export_id)
    return FileResponse(row["path"], media_type=row["content_type"], filename=row["name"])


@app.get("/api/v1/audit", tags=["audit"])
def audit_log(user: Annotated[User, Depends(current_user)]) -> dict[str, Any]:
    with store.connect() as conn:
        rows = conn.execute("SELECT actor,action,entity_id,created_at,details_json FROM audit_events WHERE tenant_id=? ORDER BY id DESC LIMIT 100", (user.tenant_id,)).fetchall()
    return {"events": [dict(row) for row in rows]}


@app.get("/api/v1/audit/download", tags=["audit"])
def download_audit_log(user: Annotated[User, Depends(current_user)]) -> JSONResponse:
    with store.connect() as conn:
        rows = conn.execute(
            "SELECT actor,action,entity_id,created_at,details_json FROM audit_events WHERE tenant_id=? ORDER BY id DESC LIMIT 1000",
            (user.tenant_id,),
        ).fetchall()
    events = []
    for row in rows:
        event = dict(row)
        event["details"] = json.loads(event.pop("details_json") or "{}")
        events.append(event)
    return JSONResponse(
        {"log_version": "matrix360-audit-log-v1", "events": events},
        headers={"Content-Disposition": 'attachment; filename="matrix360-audit-log.json"'},
    )


app.mount("/", StaticFiles(directory=WEB_ROOT, html=True), name="web")

